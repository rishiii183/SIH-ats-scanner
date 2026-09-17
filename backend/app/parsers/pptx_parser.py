import re
from typing import List, Dict, Any
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

class PPTXParser:
    @staticmethod
    def parse(file_path: str) -> List[Dict[str, Any]]:
        """
        Parses a .pptx file and returns slide-by-slide metadata & extracted content.
        """
        prs = Presentation(file_path)
        slides_data = []

        for index, slide in enumerate(prs.slides, start=1):
            slide_text_parts = []
            bullet_count = 0
            image_count = 0

            for shape in slide.shapes:
                # Extract text from shapes & text boxes
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_text_parts.append(text)
                            # Simple bullet detection check
                            if (paragraph.level > 0 or 
                                text.startswith(('•', '-', '*', '1.', '2.', '3.', '4.', '5.'))):
                                bullet_count += 1

                # Check for images / graphics / pictures
                if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.MEDIA, MSO_SHAPE_TYPE.FREEFORM):
                    image_count += 1
                elif hasattr(shape, 'image'):
                    image_count += 1

            full_text = "\n".join(slide_text_parts)
            word_count = len(re.findall(r'\b\w+\b', full_text))

            slides_data.append({
                "slide_number": index,
                "raw_text": full_text,
                "word_count": word_count,
                "bullet_count": bullet_count,
                "image_count": image_count,
            })

        return slides_data
